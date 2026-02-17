import os
import time
import random
import fitz  # PyMuPDF for PDF
import pandas as pd
from docx import Document
from pptx import Presentation
from google.api_core.exceptions import ResourceExhausted
from google.adk.agents.llm_agent import Agent


# -----------------------------
# Retry Helper
# -----------------------------
def call_with_retry(api_call, *args, **kwargs):
    """Wrapper to retry Gemini API calls with exponential backoff."""
    max_retries = 5
    for attempt in range(max_retries):
        try:
            return api_call(*args, **kwargs)
        except ResourceExhausted as e:
            retry_delay = getattr(e, "retry_delay", None)
            if retry_delay:
                wait_time = retry_delay.seconds
            else:
                wait_time = (2 ** attempt) + random.uniform(0, 1)
            print(f"[Retry] Rate limit hit. Waiting {wait_time:.2f}s before retry...")
            time.sleep(wait_time)
    raise RuntimeError("Max retries exceeded for Gemini API call")


# -----------------------------
# File Reader Utility
# -----------------------------
def read_file_content(file_path: str):
    """
    Reads the file and returns its content and extension.
    """
    ext = os.path.splitext(file_path)[1].lower()

    if ext == '.pdf':
        pdf_data = fitz.open(file_path)
        text = ""
        for page_num in range(len(pdf_data)):
            page = pdf_data[page_num]
            text += page.get_text()
        return text, ext

    elif ext == '.csv':
        df = pd.read_csv(file_path)
        return df, ext

    elif ext == '.txt':
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read(), ext

    elif ext == '.docx':
        doc = Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
        return text, ext

    elif ext == '.pptx':
        prs = Presentation(file_path)
        text = ""
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text += shape.text + "\n"
        return text, ext

    else:
        raise ValueError("Unsupported file format")


def get_file_bytes_and_mime_tool(file_path: str):
    """Tool: Return file bytes and MIME type for API upload."""
    ext = os.path.splitext(file_path)[1].lower()
    mime_map = {
        '.pdf': 'application/pdf',
        '.csv': 'text/csv',
        '.txt': 'text/plain',
        '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation',
    }
    mime_type = mime_map.get(ext, 'application/octet-stream')
    with open(file_path, 'rb') as f:
        file_bytes = f.read()
    return file_bytes, mime_type


# -----------------------------
# Specialized Agents
# -----------------------------
class CsvAnalysisAgent(Agent):
    def __init__(self):
        super().__init__(
            model='gemini-2.5-flash',
            name='csv_analysis_agent',
            description='Analyzes CSV files and answers data analysis questions.',
            instruction='Perform data analysis and answer questions based on the provided CSV file.',
            tools=[]
        )

    def analyze(self, file_bytes, mime_type, question: str):
        # Example Gemini API call with retry wrapper
        response = call_with_retry(
            self.model.file_analysis,
            file_content=file_bytes,
            mime_type=mime_type,
            question=question
        )
        return response


class PdfQnAAgent(Agent):
    def __init__(self):
        super().__init__(
            model='gemini-2.5-flash',
            name='pdf_qna_agent',
            description='Answers questions from PDF content.',
            instruction='Read the PDF and answer questions based on its content.',
            tools=[]
        )

    def answer(self, file_bytes, mime_type, question: str):
        # Example Gemini API call with retry wrapper
        response = call_with_retry(
            self.model.file_analysis,
            file_content=file_bytes,
            mime_type=mime_type,
            question=question
        )
        return response


# -----------------------------
# Orchestrator Agent
# -----------------------------
class OrchestratorAgent(Agent):
    def __init__(self, csv_agent, pdf_agent):
        super().__init__(
            model='gemini-2.5-flash',
            name='root_agent',
            description='Orchestrates specialized agents for user questions.',
            instruction='Route user questions to the appropriate agent.',
            tools=[get_file_bytes_and_mime_tool]
        )
        object.__setattr__(self, '_csv_agent', csv_agent)
        object.__setattr__(self, '_pdf_agent', pdf_agent)

    def route(self, file_path: str, question: str):
        if not file_path or not isinstance(file_path, str):
            return 'Error: No file path provided or file path is invalid.'
        if not os.path.isfile(file_path):
            return f'Error: File does not exist at path: {file_path}'

        file_bytes, mime_type = get_file_bytes_and_mime_tool(file_path)
        ext = os.path.splitext(file_path)[1].lower()

        if ext == '.csv':
            return self._csv_agent.analyze(file_bytes, mime_type, question)
        elif ext == '.pdf':
            return self._pdf_agent.answer(file_bytes, mime_type, question)
        else:
            return f'Unsupported or missing file type for analysis. Provided: {file_path}'


# -----------------------------
# Instantiate Agents
# -----------------------------
csv_agent = CsvAnalysisAgent()
pdf_agent = PdfQnAAgent()
root_agent = OrchestratorAgent(csv_agent, pdf_agent)

# Example usage:
# result = root_agent.route("sample.pdf", "Summarize the document")
# print(result)